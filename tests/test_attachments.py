import hashlib
import io

import pypdf
import pytest
from docx import Document
from helpers import (
    att_docx,
    att_pdf,
    att_png,
    att_pptx,
    att_xlsx,
    blank_pdf_bytes,
    encrypted_pdf_bytes,
    noisy_png,
    write_minimal_pdf,
    write_text_png,
)
from pptx import Presentation
from pptx.util import Inches

from emailsupersearch.ingest import attachments
from emailsupersearch.ingest.attachments import (
    INLINE_IMAGE_MIN_BYTES,
    MAX_ATTACHMENT_BYTES,
    MAX_OCR_PAGES,
    extract_text,
    kind_for,
    render_pdf_page_png,
)
from emailsupersearch.types import ParsedAttachment, ParsedMessage


def test_pdf_text_extracted():
    r = extract_text(att_pdf(write_minimal_pdf("Invoice INV-2024-00017")))
    assert "INV-2024-00017" in r.text
    assert (r.status, r.page_count, r.needs_ocr_pages) == ("ok", 1, [])


def test_scanned_pdf_pages_need_ocr():
    r = extract_text(att_pdf(blank_pdf_bytes(pages=2)))
    assert r.status == "needs_ocr" and r.needs_ocr_pages == [0, 1]
    assert r.page_count == 2


def test_pdf_with_text_and_scanned_pages_keeps_text_and_lists_scans():
    writer = pypdf.PdfWriter()
    for data in (write_minimal_pdf("Quarterly invoice INV-2024-00017"), blank_pdf_bytes(pages=1)):
        writer.append(pypdf.PdfReader(io.BytesIO(data)))
    buf = io.BytesIO()
    writer.write(buf)
    r = extract_text(att_pdf(buf.getvalue()))
    assert "INV-2024-00017" in r.text
    assert (r.status, r.page_count, r.needs_ocr_pages) == ("needs_ocr", 2, [1])


def test_ocr_page_list_is_capped():
    r = extract_text(att_pdf(blank_pdf_bytes(pages=MAX_OCR_PAGES + 3)))
    assert r.page_count == MAX_OCR_PAGES + 3
    assert r.needs_ocr_pages == list(range(MAX_OCR_PAGES))


def test_docx_xlsx_pptx_text():
    assert "quarterly" in extract_text(att_docx("quarterly figures")).text.lower()
    x = extract_text(att_xlsx([["Budget", 412000]])).text
    assert "412000" in x and "Sheet" in x
    assert "kickoff" in extract_text(att_pptx("Kickoff")).text.lower()


def test_docx_tables_headers_and_footers_are_read():
    doc = Document()
    doc.add_paragraph("body paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Vendor"
    table.rows[0].cells[1].text = "Acme"
    doc.sections[0].header.paragraphs[0].text = "Confidential header"
    doc.sections[0].footer.paragraphs[0].text = "Page footer"
    buf = io.BytesIO()
    doc.save(buf)
    att = ParsedAttachment(0, "t.docx", "", len(buf.getvalue()), "docx", data=buf.getvalue())
    r = extract_text(att)
    assert r.status == "ok"
    assert "body paragraph" in r.text and "Vendor\tAcme" in r.text
    assert "Confidential header" in r.text and "Page footer" in r.text


def test_xlsx_sheets_are_named_and_empty_cells_dropped():
    r = extract_text(att_xlsx([["Budget", None, 412000], [None, None, None], ["Total", "", 7]]))
    assert r.text.splitlines() == ["Sheet", "Budget\t412000", "Total\t7"]


def test_xlsx_cell_cap_gives_partial(monkeypatch):
    monkeypatch.setattr(attachments, "MAX_XLSX_CELLS", 4)
    r = extract_text(att_xlsx([["r1c1", "r1c2"], ["r2c1", "r2c2"], ["r3c1", "r3c2"]]))
    assert (r.status, r.reason) == ("partial", "cell_cap")
    assert "r2c2" in r.text and "r3c1" not in r.text


def test_xlsx_exactly_at_cell_cap_is_complete(monkeypatch):
    monkeypatch.setattr(attachments, "MAX_XLSX_CELLS", 4)
    r = extract_text(att_xlsx([["r1c1", "r1c2"], ["r2c1", "r2c2"]]))
    assert (r.status, r.reason) == ("ok", "") and "r2c2" in r.text


def test_pptx_tables_and_notes_are_read():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Launch"
    slide.notes_slide.notes_text_frame.text = "speaker remembers the budget"
    buf = io.BytesIO()
    prs.save(buf)
    att = ParsedAttachment(0, "d.pptx", "", len(buf.getvalue()), "pptx", data=buf.getvalue())
    r = extract_text(att)
    assert "Milestone\tLaunch" in r.text and "speaker remembers the budget" in r.text


def test_image_needs_ocr():
    r = extract_text(att_png(write_text_png("hello")))
    assert r.status == "needs_ocr"
    assert (r.page_count, r.needs_ocr_pages) == (1, [0])


def test_small_cid_image_skipped_large_cid_image_ocrd():
    small = att_png(write_text_png("logo"))
    small.content_id = "logo@x"
    skipped = extract_text(small)
    assert (skipped.status, skipped.reason, skipped.text, skipped.needs_ocr_pages) == ("skipped", "inline_image", "", [])
    big = att_png(noisy_png(INLINE_IMAGE_MIN_BYTES + 1))
    big.content_id = "scan@x"
    assert big.size > INLINE_IMAGE_MIN_BYTES and extract_text(big).status == "needs_ocr"


def test_oversized_attachment_flagged_unsupported_too_large():
    r = extract_text(
        ParsedAttachment(0, "big.pdf", "application/pdf", MAX_ATTACHMENT_BYTES + 1, "pdf", data=None)
    )
    assert (r.status, r.reason) == ("unsupported", "too_large")


def test_encrypted_pdf_flagged_unsupported():
    r = extract_text(att_pdf(encrypted_pdf_bytes()))
    assert (r.status, r.reason) == ("unsupported", "encrypted")


@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx", "pptx"])
def test_corrupt_document_fails_with_exception_class_name_instead_of_raising(kind):
    garbage = b"this is not a real document" * 20
    r = extract_text(ParsedAttachment(0, "x", "", len(garbage), kind, data=garbage))
    assert r.status == "failed" and r.reason and r.reason != "encrypted"
    assert r.text == ""


def test_document_without_data_fails_without_raising():
    r = extract_text(ParsedAttachment(0, "x.pdf", "application/pdf", 10, "pdf", data=None))
    assert (r.status, r.reason) == ("failed", "no_data")


def test_unknown_type_is_unsupported():
    r = extract_text(ParsedAttachment(0, "a.zip", "application/zip", 3, "other", data=b"abc"))
    assert (r.status, r.reason) == ("unsupported", "unsupported_type")


def _message(subject: str, body: str) -> ParsedMessage:
    return ParsedMessage(
        None, subject, "", "", [], None, None, "mtime", body, None, None, [], None, None, []
    )


@pytest.mark.parametrize("kind", ["msg", "eml"])
def test_nested_message_text_is_subject_and_body(kind):
    att = ParsedAttachment(0, "fwd", "", 5, kind, nested=_message("Re: Lease renewal", "Please sign."))
    r = extract_text(att)
    assert r.status == "ok" and r.text == "Re: Lease renewal\nPlease sign."


def test_nested_message_missing_is_unsupported():
    r = extract_text(ParsedAttachment(0, "fwd.msg", "", 5, "msg", nested=None))
    assert (r.status, r.reason) == ("unsupported", "no_nested_message")


def test_sha256_is_filled_whenever_data_is_present():
    pdf = att_pdf(encrypted_pdf_bytes())
    assert extract_text(pdf).sha256 == hashlib.sha256(pdf.data).hexdigest()  # unsupported
    ok = att_docx("hello")
    assert extract_text(ok).sha256 == hashlib.sha256(ok.data).hexdigest()
    assert extract_text(ParsedAttachment(0, "a.pdf", "", 5, "pdf", data=None)).sha256 is None


def test_result_carries_the_attachments_identity():
    att = att_docx("hello")
    att.ordinal = 3
    r = extract_text(att)
    assert (r.ordinal, r.filename, r.kind, r.size) == (3, "doc.docx", "docx", att.size)


def test_kind_for_xlsm_and_images():
    assert kind_for("a.xlsm", "") == "xlsx" and kind_for("b.webp", "") == "image"


def test_kind_for_falls_back_to_mime_when_suffix_is_missing_or_unknown():
    assert kind_for("part-0", "message/rfc822") == "eml"
    assert kind_for("mail.bin", "application/vnd.ms-outlook") == "msg"
    assert kind_for("scan", "image/gif") == "image"
    assert kind_for("download", "application/pdf") == "pdf"
    assert kind_for("notes.txt", "text/plain") == "other"
    assert kind_for("REPORT.PDF", "") == "pdf"


def test_render_pdf_page_png_returns_png():
    assert render_pdf_page_png(write_minimal_pdf("x"), 0)[:8] == b"\x89PNG\r\n\x1a\n"
