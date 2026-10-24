import io
import itertools
import re
import subprocess
import sys
import zipfile
from datetime import UTC, datetime

import openpyxl
import pypdf
import pypdfium2 as pdfium
import pytest
from helpers import (
    LONG_TEXT_3000_WORDS,
    LONG_TEXT_30000_WORDS,
    att_docx,
    att_pdf,
    att_png,
    att_pptx,
    att_xlsx,
    blank_pdf_bytes,
    encrypted_pdf_bytes,
    noisy_png,
    wait_until,
)
from PIL import Image

from emailsupersearch.diagnostics.corpus import (
    PEOPLE,
    TOPICS,
    generate,
    make_eml_bytes,
    write_minimal_pdf,
    write_text_png,
)
from emailsupersearch.ingest.eml_adapter import parse_eml

PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def pdf_text(data: bytes) -> str:
    return pdfium.PdfDocument(data)[0].get_textpage().get_text_range()


# --- generate -------------------------------------------------------------------------


def test_generate_corpus_is_deterministic(tmp_path):
    a = generate(tmp_path / "a", 4, seed=7)
    b = generate(tmp_path / "b", 4, seed=7)
    assert [p.name for p in a] == [p.name for p in b] and all(
        x.read_bytes() == y.read_bytes() for x, y in zip(a, b)
    )
    assert all(datetime.fromtimestamp(parse_eml(p).sent_ts, UTC).year == 2026 for p in a)


def test_generate_covers_all_attachment_kinds(tmp_path):
    kinds = {a.kind for p in generate(tmp_path, 4, seed=3) for a in parse_eml(p).attachments}
    assert kinds == {"pdf", "docx", "xlsx", "image"}


def test_thread_zero_root_has_invoice_number(tmp_path):
    assert "INV-2024-00017" in parse_eml(generate(tmp_path, 1, seed=1)[0]).body_text


def test_generate_names_threads_and_returns_sorted_paths(tmp_path):
    paths = generate(tmp_path / "out", 5, seed=4)
    assert paths == sorted(paths)
    assert all(re.fullmatch(r"t\d{3}-m\d{2}\.eml", p.name) for p in paths)
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == [p.name for p in paths]
    counts = {i: sum(p.name.startswith(f"t{i:03d}-") for p in paths) for i in range(5)}
    assert all(2 <= n <= 6 for n in counts.values())


def test_seed_changes_thread_lengths_only_through_the_rng(tmp_path):
    a = [p.name for p in generate(tmp_path / "a", 12, seed=1)]
    b = [p.name for p in generate(tmp_path / "b", 12, seed=2)]
    assert a != b


def test_thread_attachment_kind_cycles_and_sits_on_the_last_message(tmp_path):
    paths = generate(tmp_path, 8, seed=5)
    expected = ["pdf", "docx", "xlsx", "image"]
    for i in range(8):
        msgs = [parse_eml(p) for p in paths if p.name.startswith(f"t{i:03d}-")]
        assert [len(m.attachments) for m in msgs[:-1]] == [0] * (len(msgs) - 1)
        assert [a.kind for a in msgs[-1].attachments] == [expected[i % 4]]
    assert parse_eml(paths[-1]).attachments[0].filename.startswith("attachment-007.")


def test_dates_are_utc_2026_three_hours_apart(tmp_path):
    paths = generate(tmp_path, 3, seed=1)
    thread0 = [parse_eml(p) for p in paths if p.name.startswith("t000-")]
    assert thread0[0].sent_ts == int(datetime(2026, 1, 5, 9, 0, tzinfo=UTC).timestamp())
    assert [b.sent_ts - a.sent_ts for a, b in itertools.pairwise(thread0)] == [3 * 3600] * (len(thread0) - 1)
    assert all(datetime.fromtimestamp(parse_eml(p).sent_ts, UTC).year == 2026 for p in paths)


def test_dates_stay_in_2026_however_many_threads(tmp_path):
    paths = generate(tmp_path, 60, seed=1)
    years = {datetime.fromtimestamp(parse_eml(p).sent_ts, UTC).year for p in paths}
    assert years == {2026}
    single = tmp_path / "single.eml"
    single.write_bytes(make_eml_bytes())
    assert datetime.fromtimestamp(parse_eml(single).sent_ts, UTC).year == 2026


def test_message_ids_and_references_chain_the_thread(tmp_path):
    paths = generate(tmp_path, 1, seed=2)
    msgs = [parse_eml(p) for p in paths]
    assert [m.message_id for m in msgs] == [f"t000-m{j:02d}@corpus.example" for j in range(len(msgs))]
    assert msgs[0].in_reply_to is None and msgs[0].references == []
    for j in range(1, len(msgs)):
        assert msgs[j].in_reply_to == msgs[j - 1].message_id
        assert msgs[j].references == [m.message_id for m in msgs[:j]]
        assert msgs[j].subject == "Re: " + msgs[0].subject


def test_replies_quote_the_previous_body_and_end_with_a_signature(tmp_path):
    paths = generate(tmp_path, 1, seed=2)
    first, second = parse_eml(paths[0]), parse_eml(paths[1])
    assert first.body_text.rstrip().endswith(
        "Example Trading Ltd | +44 20 7946 0000"
    ) or first.body_text.rstrip().endswith("Example Foundation | +44 20 7946 0000")
    lines = second.body_text.splitlines()
    wrote = [i for i, line in enumerate(lines) if re.fullmatch(r"On .+, .+ wrote:", line)]
    assert len(wrote) == 1
    quoted = [line for line in lines if line.startswith("> ")]
    assert quoted and quoted[0] == "> " + first.body_text.splitlines()[0]
    assert re.search(r"Kind regards,\n.+\n.+ \| \+44 20 7946 0000\Z", second.body_text.rstrip())


def test_recipients_and_cc_rules(tmp_path):
    paths = generate(tmp_path, 2, seed=2)
    by_name = {p.name: parse_eml(p) for p in paths}
    t0_root = by_name["t000-m00.eml"]
    t1_root = by_name["t001-m00.eml"]
    assert [r.kind for r in t0_root.recipients] == ["to", "cc"]
    assert [r.kind for r in t1_root.recipients] == ["to"]
    assert all(
        [r.kind for r in m.recipients] == ["to"] for n, m in by_name.items() if not n.endswith("-m00.eml")
    )
    people = {p.addr for p in PEOPLE}
    for m in by_name.values():
        assert m.sender_addr in people and m.sender_addr not in {r.addr for r in m.recipients}
        assert {r.addr for r in m.recipients} <= people


def test_docx_and_xlsx_carry_no_wall_clock_time(tmp_path):
    seen = set()
    for p in generate(tmp_path, 3, seed=1):
        for att in parse_eml(p).attachments:
            if att.kind in {"docx", "xlsx"}:
                archive = zipfile.ZipFile(io.BytesIO(att.data))
                assert {i.date_time for i in archive.infolist()} == {(2026, 1, 1, 0, 0, 0)}
                core = archive.read("docProps/core.xml").decode()
                assert set(re.findall(r"\d{4}-\d{2}-\d{2}T[\d:]+Z", core)) == {"2026-01-01T00:00:00Z"}
                seen.add(att.kind)
    assert seen == {"docx", "xlsx"}


def test_generated_attachments_are_readable(tmp_path):
    atts = {a.kind: a for p in generate(tmp_path, 4, seed=3) for a in parse_eml(p).attachments}
    assert pdf_text(atts["pdf"].data).startswith("Invoice INV-2024-00017 for the printer lease renewal")
    assert atts["xlsx"].filename.endswith(".xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(atts["xlsx"].data))
    assert wb.sheetnames == ["Budget"] and wb["Budget"].max_row == 3
    assert atts["docx"].filename == "attachment-001.docx"
    assert atts["image"].data.startswith(PNG_MAGIC) and atts["image"].filename == "attachment-003.png"


def test_module_entry_point_writes_the_same_files(tmp_path):
    out = tmp_path / "cli"
    subprocess.run(
        [sys.executable, "-m", "emailsupersearch.diagnostics.corpus", str(out), "3", "5"],
        check=True,
        capture_output=True,
    )
    expected = generate(tmp_path / "direct", 3, seed=5)
    assert sorted(p.name for p in out.iterdir()) == [p.name for p in expected]
    assert all((out / p.name).read_bytes() == p.read_bytes() for p in expected)


# --- topics and people ----------------------------------------------------------------


def test_topics_table_shape():
    assert len(TOPICS) >= 12
    assert len({subject for subject, _, _ in TOPICS}) == len(TOPICS)
    assert TOPICS[0][1].startswith(
        "Invoice INV-2024-00017 for the printer lease renewal is attached for approval."
    )
    for _, root, replies in TOPICS:
        assert 3 <= len(re.findall(r"[.!?](?:\s|$)", root)) <= 4
        assert len(replies) == 2 and all(r.endswith((".", "?")) and len(r) < 90 for r in replies)


def test_people_span_both_example_domains():
    assert len(PEOPLE) == 6
    assert {p.addr.rsplit("@", 1)[1] for p in PEOPLE} == {"example.com", "example.org"}


# --- standalone message, PDF and PNG writers ------------------------------------------


def test_make_eml_bytes_is_a_standalone_message(tmp_path):
    assert make_eml_bytes() == make_eml_bytes()
    p = tmp_path / "one.eml"
    p.write_bytes(make_eml_bytes("w", seed=7))
    m = parse_eml(p)
    assert m.message_id == "single-7@corpus.example" and m.subject == "w"
    assert m.sender_addr == PEOPLE[7 % 6].addr and m.in_reply_to is None and m.attachments == []
    assert m.body_text.startswith(TOPICS[7 % 12][1])
    assert m.date_source == "header_date"


def test_write_minimal_pdf_text_is_extractable_with_escapes():
    text = r"Invoice INV-2024-00017 (final) paid \ done"
    data = write_minimal_pdf(text)
    assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF")
    assert pdf_text(data) == text
    assert len(pypdf.PdfReader(io.BytesIO(data)).pages) == 1


def test_write_minimal_pdf_xref_offsets_are_correct():
    data = write_minimal_pdf("x")
    xref_pos = int(data.rsplit(b"startxref", 1)[1].split()[0])
    assert data[xref_pos:].startswith(b"xref")
    entries = re.findall(rb"(\d{10}) 00000 n ", data[xref_pos:])
    for n, off in enumerate(entries, start=1):
        assert data[int(off) :].startswith(b"%d 0 obj" % n)


def test_every_topic_first_sentence_survives_a_pdf_round_trip():
    for _, root, _ in TOPICS:
        first = re.split(r"(?<=[.!?])\s+", root)[0]
        assert pdf_text(write_minimal_pdf(first)) == first


def test_write_text_png_is_black_text_on_white():
    data = write_text_png("Q3 budget approval")
    img = Image.open(io.BytesIO(data))
    assert data.startswith(PNG_MAGIC) and img.size == (1200, 400) and img.mode == "RGB"
    colours = {c for _, c in img.getcolors(maxcolors=1 << 16)}
    assert (255, 255, 255) in colours and (0, 0, 0) in colours
    assert data == write_text_png("Q3 budget approval")


def test_write_text_png_honours_size():
    assert Image.open(io.BytesIO(write_text_png("x", size=(300, 120)))).size == (300, 120)


# --- helpers module -------------------------------------------------------------------


def test_long_texts_have_exact_word_counts_and_paragraphs():
    assert len(LONG_TEXT_3000_WORDS.split()) == 3000
    assert len(LONG_TEXT_30000_WORDS.split()) == 30000
    paragraphs = LONG_TEXT_3000_WORDS.split("\n\n")
    assert len(paragraphs) > 10
    sentences = [s for s in re.split(r"(?<=\.)\s+", LONG_TEXT_3000_WORDS) if s]
    assert all(8 <= len(s.split()) <= 15 for s in sentences[:-1])
    assert all(len(para.split(". ")) <= 6 for para in paragraphs)
    assert LONG_TEXT_3000_WORDS == LONG_TEXT_3000_WORDS.strip()


def test_attachment_builders_return_parsed_attachments():
    docx = att_docx("quarterly figures\n\nsecond paragraph")
    xlsx = att_xlsx([["Budget", 412000], ["Tools", 5]])
    pptx = att_pptx("Kickoff")
    png = att_png(write_text_png("hi"))
    pdf = att_pdf(write_minimal_pdf("hello"))
    for att, kind in [(pdf, "pdf"), (docx, "docx"), (xlsx, "xlsx"), (pptx, "pptx"), (png, "image")]:
        assert att.kind == kind and att.ordinal == 0 and att.size == len(att.data)
        assert att.filename.endswith({"image": ".png"}.get(kind, f".{kind}")) and att.mime
    assert zipfile.ZipFile(io.BytesIO(docx.data)).testzip() is None
    assert zipfile.ZipFile(io.BytesIO(pptx.data)).testzip() is None
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.data)).active
    assert ws.title == "Sheet" and [c.value for c in ws[1]] == ["Budget", 412000]


def test_blank_pdf_bytes_has_empty_pages():
    doc = pdfium.PdfDocument(blank_pdf_bytes(3))
    assert len(doc) == 3 and all(doc[i].get_textpage().get_text_range() == "" for i in range(3))


def test_encrypted_pdf_bytes_needs_a_password():
    data = encrypted_pdf_bytes()
    assert pypdf.PdfReader(io.BytesIO(data)).is_encrypted
    with pytest.raises(pdfium.PdfiumError, match="password"):
        pdfium.PdfDocument(data)
    assert len(pdfium.PdfDocument(data, password="secret")) == 1


def test_noisy_png_exceeds_requested_size_and_is_deterministic():
    data = noisy_png(500_000)
    assert len(data) > 500_000 and data.startswith(PNG_MAGIC) and data == noisy_png(500_000)
    assert len(noisy_png(1)) > 1


def test_wait_until_returns_truthy_value_or_fails():
    calls = iter([0, "", "ready"])
    assert wait_until(lambda: next(calls), 2, interval=0.01) == "ready"
    with pytest.raises(pytest.fail.Exception, match="not met within 0.1s"):
        wait_until(lambda: False, 0.1, interval=0.02)


def test_corpus_dir_fixture(corpus_dir):
    names = sorted(p.name for p in corpus_dir.iterdir())
    assert names == sorted(p.name for p in corpus_dir.glob("*.eml")) and "t005-m00.eml" in names
    assert corpus_dir.name == "corpus"
