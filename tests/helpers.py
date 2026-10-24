"""Builders shared by the test modules: sample documents, long texts and a polling helper."""

from __future__ import annotations

import io
import random
import time
from collections.abc import Callable, Sequence
from email.message import EmailMessage
from typing import Any

import openpyxl
import pypdf
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from emailsupersearch.diagnostics.corpus import (
    make_eml_bytes,
    write_minimal_pdf,
    write_text_png,
)
from emailsupersearch.types import ParsedAttachment

__all__ = [
    "LONG_TEXT_3000_WORDS",
    "LONG_TEXT_30000_WORDS",
    "att_docx",
    "att_pdf",
    "att_png",
    "att_pptx",
    "att_xlsx",
    "blank_pdf_bytes",
    "encrypted_pdf_bytes",
    "make_eml_bytes",
    "make_eml_with_png_bytes",
    "noisy_png",
    "wait_until",
    "write_minimal_pdf",
    "write_text_png",
]

# fmt: off
_WORDS = [
    "account", "agenda", "approval", "archive", "audit", "backup", "balance", "budget", "calendar",
    "capacity", "certificate", "client", "contract", "cost", "deadline", "delivery", "department",
    "invoice", "laptop", "lease", "ledger", "mailbox", "meeting", "migration", "network", "notice",
    "office", "order", "outage", "package", "payment", "permit", "planning", "policy", "printer",
    "procurement", "project", "quarter", "quote", "receipt", "record", "renewal", "report",
    "request", "review", "schedule", "security", "server", "service", "signature", "spreadsheet",
    "supplier", "support", "template", "ticket", "training", "update", "vendor", "warranty",
    "workshop",
]
# fmt: on
assert len(_WORDS) == 60


def _long_text(n_words: int) -> str:
    """Seeded filler prose: sentences of 8 to 15 words, a blank line every 6 sentences, cut to ``n_words``."""
    rng = random.Random(11)
    out: list[str] = []
    words = 0
    sentence = 0
    while words < n_words:
        length = min(rng.randint(8, 15), n_words - words)
        picked = [rng.choice(_WORDS) for _ in range(length)]
        picked[0] = picked[0].capitalize()
        picked[-1] += "."
        if sentence:
            out.append("\n\n" if sentence % 6 == 0 else " ")
        out.append(" ".join(picked))
        words += length
        sentence += 1
    return "".join(out)


LONG_TEXT_3000_WORDS = _long_text(3000)
LONG_TEXT_30000_WORDS = _long_text(30000)


def _attachment(filename: str, mime: str, kind: str, data: bytes) -> ParsedAttachment:
    return ParsedAttachment(0, filename, mime, len(data), kind, data=data)


def att_pdf(data: bytes) -> ParsedAttachment:
    return _attachment("doc.pdf", "application/pdf", "pdf", data)


def att_docx(text: str) -> ParsedAttachment:
    """A Word file with one paragraph per blank-line-separated block of ``text``."""
    doc = Document()
    for paragraph in text.split("\n\n"):
        doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return _attachment(
        "doc.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
        buf.getvalue(),
    )


def att_xlsx(rows: Sequence[Sequence[Any]]) -> ParsedAttachment:
    """A workbook with one sheet (openpyxl's default name ``Sheet``) holding ``rows``."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return _attachment(
        "book.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
        buf.getvalue(),
    )


def att_pptx(text: str) -> ParsedAttachment:
    """A one-slide deck with ``text`` in a text box."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return _attachment(
        "deck.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "pptx",
        buf.getvalue(),
    )


def att_png(data: bytes) -> ParsedAttachment:
    return _attachment("image.png", "image/png", "image", data)


def make_eml_with_png_bytes(text: str) -> bytes:
    """A message with a plain-text body and ``text`` rendered as ``scan.png`` (``image/png``, no Content-ID)."""
    msg = EmailMessage()
    msg["Subject"] = "scanned page"
    msg["From"] = "sender@example.com"
    msg["To"] = "recipient@example.org"
    msg["Message-ID"] = "<scan-1@example.com>"
    msg.set_content("The scanned page is attached.")
    msg.add_attachment(write_text_png(text), maintype="image", subtype="png", filename="scan.png")
    return msg.as_bytes()


def blank_pdf_bytes(pages: int) -> bytes:
    """A PDF of ``pages`` empty pages: no text layer, like a scan with the image removed."""
    return write_minimal_pdf("", pages=pages)


def noisy_png(min_bytes: int) -> bytes:
    """A PNG larger than ``min_bytes``: seeded RGB noise (it barely compresses) with ``scan`` on top."""
    width, height = 400, 300
    font = ImageFont.load_default(size=60)
    while True:
        noise = random.Random(1).randbytes(width * height * 3)
        image = Image.frombytes("RGB", (width, height), noise)
        ImageDraw.Draw(image).text(
            (20, 20), "scan", fill="white", font=font, stroke_width=3, stroke_fill="black"
        )
        buf = io.BytesIO()
        image.save(buf, "PNG")
        if buf.tell() > min_bytes:
            return buf.getvalue()
        width, height = width * 2, height * 2


def encrypted_pdf_bytes() -> bytes:
    """A one-page PDF that needs the password ``secret`` to open."""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(200, 100)
    writer.encrypt("secret")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def wait_until(pred: Callable[[], Any], timeout_s: float, interval: float = 0.1) -> Any:
    """Poll ``pred`` until it returns something truthy, which is returned; fail the test on timeout."""
    deadline = time.monotonic() + timeout_s
    while True:
        value = pred()
        if value:
            return value
        if time.monotonic() >= deadline:
            pytest.fail(f"condition not met within {timeout_s}s")
        time.sleep(interval)
