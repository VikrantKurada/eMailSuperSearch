"""Excel export: columns, links, row cap, text safety and atomic save."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from emailsupersearch.export.excel import EXPORT_MAX_ROWS, export_results
from emailsupersearch.search.results import AttachmentInfo, MessageCard
from emailsupersearch.types import Person

SENT = 1_773_705_600  # 2026-03-17 00:00 UTC

HEADERS = [
    "Date",
    "From",
    "To",
    "CC",
    "Subject",
    "Chain",
    "Attachments",
    "Preview",
    "Open",
    "Path",
]


def _card(key="k1", subject="Quarterly plan", paths=(), **overrides):
    values = {
        "message_key": key,
        "subject": subject,
        "sent_ts": SENT,
        "sender": Person("Ada Lovelace", "ada@example.com"),
        "to": [Person("Bob", "bob@example.com")],
        "cc": [Person("", "cc@example.com")],
        "snippet": "the plan is attached",
        "attachments": [],
        "paths": list(paths),
        "chain_id": None,
        "score": 1.0,
    }
    values.update(overrides)
    return MessageCard(**values)


def _attachment(filename):
    return AttachmentInfo(1, filename, "pdf", 10, "ok", None, False, 1)


def test_export_writes_expected_columns_and_hyperlink(tmp_path):
    msg_path = tmp_path / "mail" / "plan.msg"
    card = _card(
        paths=[str(msg_path)],
        attachments=[_attachment("plan.pdf"), _attachment("notes.docx")],
        chain_id="chain-1",
    )

    info = export_results([card], tmp_path / "out" / "results.xlsx")

    assert info.path == tmp_path / "out" / "results.xlsx"
    assert info.rows == 1
    assert info.truncated is False
    ws = load_workbook(info.path).active
    assert [c.value for c in ws[1]] == HEADERS
    assert all(c.font.bold for c in ws[1])
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions
    assert ws["A2"].value == "2026-03-17 00:00"
    assert ws["B2"].value == "Ada Lovelace <ada@example.com>"
    assert ws["C2"].value == "Bob <bob@example.com>"
    assert ws["D2"].value == "cc@example.com"
    assert ws["E2"].value == "Quarterly plan"
    assert ws["F2"].value == "chain-1"
    assert ws["G2"].value == "plan.pdf; notes.docx"
    assert ws["H2"].value == "the plan is attached"
    assert ws["H2"].alignment.wrap_text
    assert ws.column_dimensions["H"].width == 60
    assert ws["I2"].value == "Open"
    assert ws["I2"].hyperlink.target.startswith("file:///")
    assert ws["I2"].style == "Hyperlink"
    assert ws["J2"].value == str(msg_path)
    assert ws.column_dimensions["E"].width == len("Quarterly plan") + 2


def test_export_caps_rows_and_reports_truncation(tmp_path):
    cards = [_card(key=f"k{i}", subject=f"s{i}", snippet="", sent_ts=None) for i in range(6000)]

    info = export_results(cards, tmp_path / "big.xlsx")

    assert EXPORT_MAX_ROWS == 5000
    assert info.rows == 5000
    assert info.truncated is True
    ws = load_workbook(info.path).active
    assert ws.max_row == 5001
    assert ws["A2"].value in (None, "")


def test_export_percent_encodes_path(tmp_path):
    odd = tmp_path / "my mail" / "a#b.msg"

    info = export_results([_card(paths=[str(odd)])], tmp_path / "out.xlsx")

    target = load_workbook(info.path).active["I2"].hyperlink.target
    assert "%20" in target
    assert "%23" in target


def test_formula_like_text_is_stored_as_text(tmp_path):
    evil = '=HYPERLINK("http://evil")'
    card = _card(subject=evil, snippet="@SUM(A1)", attachments=[_attachment("-cmd.exe")])

    info = export_results([card], tmp_path / "out.xlsx")

    ws = load_workbook(info.path).active
    assert ws["E2"].data_type == "s"
    assert ws["E2"].value == evil
    assert ws["H2"].data_type == "s"
    assert ws["H2"].value == "@SUM(A1)"
    assert ws["G2"].data_type == "s"
    assert ws["G2"].value == "-cmd.exe"


def test_chain_subjects_mapping_used(tmp_path):
    cards = [
        _card(key="a", chain_id="c1"),
        _card(key="b", chain_id="c2"),
        _card(key="c", chain_id=None),
    ]

    info = export_results(cards, tmp_path / "out.xlsx", chain_subjects={"c1": "Budget thread"})

    ws = load_workbook(info.path).active
    assert ws["F2"].value == "Budget thread"
    assert ws["F3"].value == "c2"
    assert ws["F4"].value in (None, "")


def test_relative_path_is_resolved_for_the_link(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    info = export_results([_card(paths=["mail/rel.msg"])], tmp_path / "out.xlsx")

    target = load_workbook(info.path).active["I2"].hyperlink.target
    assert target.startswith("file:///")
    assert target.endswith("/mail/rel.msg")


def test_card_without_path_has_empty_open_cell(tmp_path):
    info = export_results([_card(paths=[])], tmp_path / "out.xlsx")

    ws = load_workbook(info.path).active
    assert ws["I2"].value in (None, "")
    assert ws["I2"].hyperlink is None
    assert ws["J2"].value in (None, "")


def test_export_leaves_no_temp_file_behind(tmp_path):
    export_results([_card()], tmp_path / "out.xlsx")

    assert sorted(p.name for p in Path(tmp_path).iterdir()) == ["out.xlsx"]


def test_control_characters_are_dropped_not_fatal(tmp_path):
    card = _card(subject="bad\x00subject\x1f", snippet="line1\nline2\ttab")

    info = export_results([card], tmp_path / "out.xlsx")

    ws = load_workbook(info.path).active
    assert ws["E2"].value == "badsubject"
    assert ws["H2"].value == "line1\nline2\ttab"


def test_failing_save_propagates_and_leaves_no_temp_file(tmp_path, monkeypatch):
    def boom(self, filename):
        raise RuntimeError("boom")

    monkeypatch.setattr(Workbook, "save", boom)

    with pytest.raises(RuntimeError, match="boom"):
        export_results([_card()], tmp_path / "out.xlsx")

    assert list(tmp_path.iterdir()) == []


def test_failing_rename_keeps_old_file_and_removes_temp(tmp_path, monkeypatch):
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"old")

    def nope(src, dst):
        raise OSError("nope")

    monkeypatch.setattr("emailsupersearch.export.excel.os.replace", nope)

    with pytest.raises(OSError, match="nope"):
        export_results([_card()], out)

    assert out.read_bytes() == b"old"
    assert [p.name for p in tmp_path.iterdir()] == ["out.xlsx"]


def test_cleanup_failure_does_not_mask_the_original_error(tmp_path, monkeypatch):
    def nope(src, dst):
        raise OSError("nope")

    def locked(self, missing_ok=False):
        raise PermissionError("locked")

    monkeypatch.setattr("emailsupersearch.export.excel.os.replace", nope)
    monkeypatch.setattr(Path, "unlink", locked)

    with pytest.raises(OSError, match="nope"):
        export_results([_card()], tmp_path / "out.xlsx")


def test_lone_surrogate_text_is_exported_with_replacement_character(tmp_path):
    bad = "bad" + chr(0xDC80) + "byte"
    card = _card(subject=bad, snippet=bad, attachments=[_attachment(bad)])

    info = export_results([card], tmp_path / "out.xlsx")

    ws = load_workbook(info.path).active
    expected = "bad" + chr(0xFFFD) + "byte"
    assert ws["E2"].value == expected
    assert ws["G2"].value == expected
    assert ws["H2"].value == expected


@pytest.mark.parametrize("sent_ts", [10**13, -(10**13), 2**70])
def test_out_of_range_timestamp_leaves_date_empty(tmp_path, sent_ts):
    info = export_results([_card(sent_ts=sent_ts)], tmp_path / "out.xlsx")

    ws = load_workbook(info.path).active
    assert info.rows == 1
    assert ws["A2"].value in (None, "")
    assert ws["E2"].value == "Quarterly plan"


def test_negative_max_rows_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="max_rows"):
        export_results([_card()], tmp_path / "out.xlsx", max_rows=-1)

    assert list(tmp_path.iterdir()) == []


def test_zero_max_rows_writes_header_only(tmp_path):
    info = export_results([_card(), _card(key="k2")], tmp_path / "out.xlsx", max_rows=0)

    ws = load_workbook(info.path).active
    assert info.rows == 0
    assert info.truncated is True
    assert ws.max_row == 1
    assert ws["A1"].value == "Date"


def test_empty_chain_subject_falls_back_to_chain_id(tmp_path):
    info = export_results(
        [_card(chain_id="c1")], tmp_path / "out.xlsx", chain_subjects={"c1": ""}
    )

    assert load_workbook(info.path).active["F2"].value == "c1"


def test_text_longer_than_excel_cell_limit_is_clamped(tmp_path):
    info = export_results([_card(snippet="x" * 40_000)], tmp_path / "out.xlsx")

    assert load_workbook(info.path).active["H2"].value == "x" * 32_767
