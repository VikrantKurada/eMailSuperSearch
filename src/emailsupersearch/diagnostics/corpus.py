"""Deterministic synthetic mail corpus.

Ships in the wheel because ``bench`` runs on the owner's laptop, and every ingest test
parses it. Nothing here reads the clock or the environment, so one ``(n_threads, seed)``
pair always yields byte-identical files.

Run ``python -m emailsupersearch.diagnostics.corpus OUT_DIR N_THREADS [SEED]`` to write one.
"""

from __future__ import annotations

import io
import random
import re
import sys
import textwrap
import zipfile
from datetime import UTC, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime, formataddr
from pathlib import Path

import openpyxl
from docx import Document
from openpyxl.writer.excel import ExcelWriter
from PIL import Image, ImageDraw, ImageFont

from emailsupersearch.types import Person

TOPICS: list[tuple[str, str, list[str]]] = [
    (
        "Printer lease renewal",
        (
            "Invoice INV-2024-00017 for the printer lease renewal is attached for approval. "
            "The lease covers the three floor printers at the London office for another 24 months. "
            "The monthly charge rises by four percent compared with the current agreement. "
            "Please confirm by Friday so that finance can release payment."
        ),
        [
            "Approved from my side, please pass the invoice to finance.",
            "I have asked the supplier to confirm the toner allowance before we pay.",
        ],
    ),
    (
        "Q3 budget approval",
        (
            "The Q3 budget approval pack is ready for your review. "
            "It sets payroll spend at 1,200,000 pounds and keeps travel flat against the previous quarter. "
            "Marketing asks for an extra 40,000 pounds for the autumn campaign. "
            "Sign-off is needed before the board meets on the 14th."
        ),
        [
            "I can approve everything except the extra marketing spend.",
            "Finance has updated the forecast with the revised payroll figures.",
        ],
    ),
    (
        "Office move to the 4th floor",
        (
            "The office move to the 4th floor is scheduled for the second weekend of March. "
            "Each team will receive labelled crates and a seating plan by Monday. "
            "The facilities desk will reconnect monitors and docking stations before staff arrive. "
            "Please clear your lockers and tell us about any ergonomic chair needs."
        ),
        [
            "Our team will pack on Friday afternoon and label the crates.",
            "The seating plan still places two of us by the lift, can we swap?",
        ],
    ),
    (
        "VPN outage on Tuesday morning",
        (
            "The VPN gateway went down on Tuesday morning and remote staff could not reach the file servers. "
            "The network team traced the fault to an expired certificate on the firewall. "
            "Service was restored at 10:40 after the certificate was renewed. "
            "We will add an expiry alert so this does not happen again."
        ),
        [
            "Thanks for the quick fix, I could reconnect at lunchtime.",
            "Can we get a short incident report for the service review?",
        ],
    ),
    (
        "New starter laptop request",
        (
            "A new starter joins the analytics team on Monday and needs a laptop. "
            "The standard build includes 16 gigabytes of memory, a docking station and two monitors. "
            "Please raise the order with the IT service desk and quote the cost centre. "
            "Delivery usually takes five working days."
        ),
        [
            "The order is placed and the laptop should arrive on Thursday.",
            "I will collect it from reception and set up the accounts.",
        ],
    ),
    (
        "Vendor contract renewal for Northwind Logistics",
        (
            "The courier contract with Northwind Logistics expires at the end of next month. "
            "Their renewal quote is eight percent higher than the present rate card. "
            "Procurement suggests asking two other carriers for comparable offers. "
            "We need a decision on the notice period before the end of the week."
        ),
        [
            "Legal says we must give sixty days notice to avoid auto renewal.",
            "I have requested quotes from two rival couriers.",
        ],
    ),
    (
        "Team offsite planning",
        (
            "The team offsite is planned for the first week of June at the lakeside conference centre. "
            "The agenda has two workshop mornings, a planning afternoon and an evening dinner. "
            "Travel is by coach from the office and accommodation is booked for two nights. "
            "Please send your dietary requirements to the organisers."
        ),
        [
            "Count me in, but I can only attend the first day.",
            "The dinner venue needs final numbers by the end of next week.",
        ],
    ),
    (
        "Mandatory security awareness training",
        (
            "Mandatory security awareness training must be completed by every employee this quarter. "
            "The course takes forty minutes and covers phishing, password hygiene and removable media. "
            "Completion is tracked in the learning portal and reported to the compliance team. "
            "Managers should remind anyone who has not started."
        ),
        [
            "I finished the phishing module and passed the quiz.",
            "Two contractors cannot log in to the learning portal yet.",
        ],
    ),
    (
        "Parking permits for the new car park",
        (
            "Parking permits for the new car park will be issued from the first of next month. "
            "Spaces are allocated by team and are limited to one permit per employee. "
            "Cyclists can apply for a locker and shower pass instead. "
            "Fill in the permit form and include your vehicle registration number."
        ),
        [
            "I have submitted the form with my registration number.",
            "Is there a visitor allowance for client days?",
        ],
    ),
    (
        "Quarterly review slides",
        (
            "The slides for the quarterly review need your latest numbers by Wednesday noon. "
            "Each team gets one slide covering results, risks and next steps. "
            "Please use the shared template so the deck keeps a consistent look. "
            "The leadership team will present it on Thursday afternoon."
        ),
        [
            "My slide is uploaded and the risk section is complete.",
            "Could we move the revenue chart to the front of the deck?",
        ],
    ),
    (
        "Expense policy update",
        (
            "The expense policy update takes effect from the first of April. "
            "Meals on business trips are now capped at forty pounds per day and rail travel must be booked in standard class. "
            "Receipts have to be uploaded within thirty days or the claim is rejected. "
            "Questions can go to the finance helpdesk."
        ),
        [
            "Does the new meal cap apply to client dinners as well?",
            "I have updated my team about the receipt deadline.",
        ],
    ),
    (
        "Data retention audit",
        (
            "The annual data retention audit begins next week and covers shared drives, mailboxes and backup tapes. "
            "Each department must list the records it keeps and the legal reason for keeping them. "
            "Anything past its retention period will be scheduled for secure deletion. "
            "The auditors need the inventory spreadsheets by the end of the month."
        ),
        [
            "Our inventory is nearly done, only the archive boxes are missing.",
            "Please confirm whether the 2019 contracts fall under the audit.",
        ],
    ),
]

PEOPLE: list[Person] = [
    Person("Alice Morgan", "alice.morgan@example.com"),
    Person("Ben Carter", "ben.carter@example.com"),
    Person("Chloe Davies", "chloe.davies@example.com"),
    Person("Dev Patel", "dev.patel@example.org"),
    Person("Elena Rossi", "elena.rossi@example.org"),
    Person("Farid Khan", "farid.khan@example.org"),
]

_ORGANISATION = {"example.com": "Example Trading Ltd", "example.org": "Example Foundation"}
_FIRST_SENT = datetime(2026, 1, 5, 9, 0, tzinfo=UTC)
_FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
_FIXED_DOC_TIME = datetime(2026, 1, 1)
_ATTACHMENT_KINDS = ["pdf", "docx", "xlsx", "png"]
_ATTACHMENT_MIME = {
    "pdf": ("application", "pdf"),
    "docx": ("application", "vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "xlsx": ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "png": ("image", "png"),
}


def _sent_at(thread: int, index: int) -> datetime:
    # Weekly per thread, wrapping at 350 days so every Date stays in 2026 however many threads.
    return _FIRST_SENT + timedelta(days=(7 * thread) % 350, hours=3 * index)


def _signature(person: Person) -> str:
    organisation = _ORGANISATION[person.addr.rsplit("@", 1)[1]]
    return f"Kind regards,\n{person.name}\n{organisation} | +44 20 7946 0000"


def _quote(text: str) -> str:
    return "\n".join(("> " + line).rstrip() for line in text.splitlines())


def _compose(text: str, author: Person, quoted: tuple[Person, datetime, str] | None) -> str:
    """A message body: the quoted previous message under its attribution line, the new text, a signature."""
    blocks = []
    if quoted is not None:
        who, when, previous = quoted
        blocks.append(f"On {format_datetime(when)}, {who.name} wrote:\n{_quote(previous)}")
    blocks.append(text)
    blocks.append(_signature(author))
    return "\n\n".join(blocks)


def _build_eml(
    *,
    subject: str,
    sender: Person,
    to: Person,
    cc: Person | None,
    sent: datetime,
    message_id: str,
    in_reply_to: str | None,
    references: list[str],
    body: str,
    attachment: tuple[str, bytes] | None,
) -> bytes:
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = formataddr((sender.name, sender.addr))
    msg["To"] = formataddr((to.name, to.addr))
    if cc is not None:
        msg["Cc"] = formataddr((cc.name, cc.addr))
    msg["Date"] = format_datetime(sent)
    msg["Message-ID"] = message_id
    if in_reply_to is not None:
        msg["In-Reply-To"] = in_reply_to
        msg["References"] = " ".join(references)
    msg.set_content(body)
    if attachment is not None:
        filename, data = attachment
        maintype, subtype = _ATTACHMENT_MIME[filename.rsplit(".", 1)[1]]
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
        # The generator would otherwise pick a random multipart boundary.
        msg.set_boundary("corpus-" + message_id.strip("<>").split("@")[0])
    return msg.as_bytes()


def _first_sentence(text: str) -> str:
    return re.split(r"(?<=[.!?])\s+", text, maxsplit=1)[0]


def _fix_zip_times(data: bytes) -> bytes:
    """Rewrite an OOXML zip so every entry carries the same timestamp."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.create_system = 0
            fixed.external_attr = 0o600 << 16
            dst.writestr(fixed, src.read(info))
    return out.getvalue()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = _FIXED_DOC_TIME
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return _fix_zip_times(buf.getvalue())


def _xlsx_bytes(thread: int, subject: str) -> bytes:
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = _FIXED_DOC_TIME
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Amount (GBP)", "Owner"])
    ws.append([subject, 1000 + 250 * thread, PEOPLE[thread % len(PEOPLE)].name])
    ws.append(["Contingency", 150 + 10 * thread, PEOPLE[(thread + 1) % len(PEOPLE)].name])
    buf = io.BytesIO()
    # ``wb.save`` stamps ``modified`` with the clock; the writer it wraps does not.
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        ExcelWriter(wb, archive).save()
    return _fix_zip_times(buf.getvalue())


def _attachment_for_thread(thread: int) -> tuple[str, bytes]:
    subject, root, _ = TOPICS[thread % len(TOPICS)]
    kind = _ATTACHMENT_KINDS[thread % len(_ATTACHMENT_KINDS)]
    if kind == "pdf":
        data = write_minimal_pdf(_first_sentence(root))
    elif kind == "docx":
        data = _docx_bytes(re.split(r"(?<=[.!?])\s+", root))
    elif kind == "xlsx":
        data = _xlsx_bytes(thread, subject)
    else:
        data = write_text_png(subject)
    return f"attachment-{thread:03d}.{kind}", data


def generate(out_dir: Path, n_threads: int, seed: int) -> list[Path]:
    """Write ``n_threads`` reply chains of 2 to 6 messages each and return the files, sorted."""
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    paths: list[Path] = []
    for thread in range(n_threads):
        subject, root, replies = TOPICS[thread % len(TOPICS)]
        count = rng.randint(2, 6)
        ids: list[str] = []
        previous: tuple[Person, datetime, str] | None = None
        for index in range(count):
            sender = PEOPLE[(thread + index) % len(PEOPLE)]
            sent = _sent_at(thread, index)
            message_id = f"<t{thread:03d}-m{index:02d}@corpus.example>"
            text = root if index == 0 else replies[(index - 1) % len(replies)]
            body = _compose(text, sender, previous)
            data = _build_eml(
                subject=subject if index == 0 else f"Re: {subject}",
                sender=sender,
                to=PEOPLE[(thread + index + 1) % len(PEOPLE)],
                cc=PEOPLE[(thread + 3) % len(PEOPLE)] if index == 0 and thread % 2 == 0 else None,
                sent=sent,
                message_id=message_id,
                in_reply_to=ids[-1] if ids else None,
                references=list(ids),
                body=body,
                attachment=_attachment_for_thread(thread) if index == count - 1 else None,
            )
            path = out_dir / f"t{thread:03d}-m{index:02d}.eml"
            path.write_bytes(data)
            paths.append(path)
            ids.append(message_id)
            previous = (sender, sent, body)
    return sorted(paths)


def make_eml_bytes(subject: str = "new", seed: int = 99) -> bytes:
    """One standalone message, for tests that drop a new file into a watched folder."""
    sender = PEOPLE[seed % len(PEOPLE)]
    body = _compose(TOPICS[seed % len(TOPICS)][1], sender, None)
    return _build_eml(
        subject=subject,
        sender=sender,
        to=PEOPLE[(seed + 1) % len(PEOPLE)],
        cc=None,
        sent=_sent_at(seed, 0),
        message_id=f"<single-{seed}@corpus.example>",
        in_reply_to=None,
        references=[],
        body=body,
        attachment=None,
    )


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_minimal_pdf(text: str, pages: int = 1) -> bytes:
    """A hand-written PDF whose first page holds ``text`` in one text object.

    With empty ``text`` every page is blank, which is how scanned (image-only) pages look.
    """
    contents = [b""] * pages
    if pages and text:
        contents[0] = f"BT /F1 12 Tf 72 720 Td ({_pdf_escape(text)}) Tj ET".encode("latin-1", "replace")
    kids = " ".join(f"{4 + 2 * k} 0 R" for k in range(pages))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for k, content in enumerate(contents):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * k} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n" % (len(objects) + 1, xref_at)
    out += b"%%EOF\n"
    return bytes(out)


def write_text_png(text: str, size: tuple[int, int] = (1200, 400)) -> bytes:
    """Black text on a white RGB image, wrapped every 50 characters."""
    image = Image.new("RGB", size, "white")
    font = ImageFont.load_default(size=40)
    ImageDraw.Draw(image).multiline_text(
        (40, 40), "\n".join(textwrap.wrap(text, 50)), fill="black", font=font
    )
    buf = io.BytesIO()
    image.save(buf, "PNG")
    return buf.getvalue()


if __name__ == "__main__":
    generate(Path(sys.argv[1]), int(sys.argv[2]), int(sys.argv[3]) if len(sys.argv) > 3 else 1)
