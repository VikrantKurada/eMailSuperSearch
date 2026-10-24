"""Write search results to an Excel workbook, one row per message with a link to the original file."""

from __future__ import annotations

import contextlib
import io
import os
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from emailsupersearch.search.results import MessageCard
from emailsupersearch.types import Person

EXPORT_MAX_ROWS = 5000

_HEADERS = (
    "Date",
    "From",
    "To",
    "CC",
    "Subject",
    "Chain",
    "Attachments",
    "Preview",
    "Open",
    "Path",
)
_PREVIEW_COLUMN = _HEADERS.index("Preview") + 1
_OPEN_COLUMN = _HEADERS.index("Open") + 1
_MAX_WIDTH = 60
_MAX_CELL_CHARS = 32_767  # Excel refuses to open a workbook whose cell text is longer
_REPLACEMENT_CHAR = chr(0xFFFD)
_LONE_SURROGATE_RE = re.compile(f"[{chr(0xD800)}-{chr(0xDFFF)}]")


@dataclass
class ExportInfo:
    """What was written: the file, the number of message rows and whether the list was cut at the cap."""

    path: Path
    rows: int
    truncated: bool


def _write_text(cell: Cell, value: str) -> None:
    """Store ``value`` as a string even when it starts with ``=``, so Excel never evaluates it.

    Text that a worksheet cannot hold is repaired rather than failing the export: lone surrogates (what
    undecodable bytes become in mail text) turn into U+FFFD, control characters openpyxl rejects are
    dropped, and the text is cut at Excel's cell limit.
    """
    value = _LONE_SURROGATE_RE.sub(_REPLACEMENT_CHAR, value)
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:_MAX_CELL_CHARS]
    cell.data_type = "s"


def _person(person: Person) -> str:
    return f"{person.name} <{person.addr}>" if person.name else person.addr


def _people(people: list[Person]) -> str:
    return "; ".join(_person(p) for p in people)


def _row_values(card: MessageCard, chain_subjects: dict[str, str] | None) -> list[str]:
    """The cells of one card from Date to Preview, in column order."""
    when = ""
    if card.sent_ts is not None:
        # A timestamp the platform cannot turn into a date leaves the cell empty.
        with contextlib.suppress(OverflowError, OSError, ValueError):
            when = datetime.fromtimestamp(card.sent_ts, UTC).strftime("%Y-%m-%d %H:%M")
    chain = card.chain_id or ""
    if chain:
        chain = (chain_subjects or {}).get(chain) or chain
    return [
        when,
        _person(card.sender),
        _people(card.to),
        _people(card.cc),
        card.subject,
        chain,
        "; ".join(a.filename for a in card.attachments),
        card.snippet,
    ]


def _file_uri(path: str) -> str:
    target = Path(path)
    if not target.is_absolute():
        target = target.resolve()
    return target.as_uri()


def export_results(
    messages: list[MessageCard],
    out_path: Path,
    max_rows: int = EXPORT_MAX_ROWS,
    *,
    chain_subjects: dict[str, str] | None = None,
) -> ExportInfo:
    """Write up to ``max_rows`` cards to ``out_path`` and say how many rows went in and if the rest was cut.

    The Chain column shows ``chain_subjects[chain_id]`` when it is given and not empty, else the chain id.
    The workbook is saved beside ``out_path`` first and renamed, so a half-written file never sits at
    ``out_path``. A negative ``max_rows`` raises ``ValueError``; zero writes the header only.
    """
    if max_rows < 0:
        raise ValueError(f"max_rows must not be negative, got {max_rows}")
    out_path = Path(out_path)
    kept = messages[:max_rows]
    truncated = len(messages) > max_rows

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    widths = [len(h) for h in _HEADERS]
    bold = Font(bold=True)
    for col, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        _write_text(cell, header)
        cell.font = bold

    for row, card in enumerate(kept, start=2):
        path = card.paths[0] if card.paths else ""
        row_values = [*_row_values(card, chain_subjects), "Open" if path else "", path]
        for col, value in enumerate(row_values, start=1):
            if value:
                _write_text(ws.cell(row=row, column=col), value)
                widths[col - 1] = max(widths[col - 1], len(value))
        ws.cell(row=row, column=_PREVIEW_COLUMN).alignment = Alignment(wrap_text=True)
        if path:
            link = ws.cell(row=row, column=_OPEN_COLUMN)
            link.hyperlink = _file_uri(path)
            link.style = "Hyperlink"

    for col, longest in enumerate(widths, start=1):
        width = _MAX_WIDTH if col == _PREVIEW_COLUMN else min(_MAX_WIDTH, longest + 2)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Serialise in memory so a failing save never leaves openpyxl holding the temp file open.
    buffer = io.BytesIO()
    wb.save(buffer)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    temp = out_path.with_name(f"{out_path.name}.{os.getpid()}.tmp")
    try:
        temp.write_bytes(buffer.getvalue())
        os.replace(temp, out_path)
    except BaseException:
        with contextlib.suppress(OSError):
            temp.unlink(missing_ok=True)
        raise
    return ExportInfo(path=out_path, rows=len(kept), truncated=truncated)
