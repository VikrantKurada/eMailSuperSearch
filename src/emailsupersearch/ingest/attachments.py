"""Turn an attachment into searchable text, or flag it, without ever failing the email it belongs to.

``extract_text`` is the single entry point and never raises: a file over the size cap, an encrypted PDF,
a corrupt Office document or a type we do not read each come back as an ``AttachmentText`` whose
``status`` and ``reason`` say what happened. Scanned PDF pages and images are only *flagged*
(``needs_ocr`` plus the page numbers); the OCR step reads them later, through ``render_pdf_page_png``.
"""

from __future__ import annotations

import hashlib
import io
import os
from collections.abc import Callable
from typing import NamedTuple

import docx
import openpyxl
import pptx
import pypdfium2
from pptx.shapes.group import GroupShape

from emailsupersearch.types import (
    AttachmentKind,
    AttachmentStatus,
    AttachmentText,
    ParsedAttachment,
)

MAX_ATTACHMENT_BYTES = 25 * 1024 * 1024  # larger attachments are not opened at all
INLINE_IMAGE_MIN_BYTES = 100 * 1024  # a smaller image with a Content-ID is a logo or signature, not content
LOW_TEXT_CHARS_PER_PAGE = 20  # a PDF page with fewer characters than this is treated as a scan
MAX_OCR_PAGES = 50  # per PDF, so one huge scan cannot dominate the OCR queue
MAX_XLSX_CELLS = 50_000  # per workbook, empty cells included, so a runaway sheet cannot stall indexing
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

_SUFFIX_KINDS: dict[str, AttachmentKind] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".pptx": "pptx",
    **dict.fromkeys(IMAGE_SUFFIXES, "image"),
    ".msg": "msg",
    ".eml": "eml",
}
_DATA_KINDS = frozenset({"pdf", "docx", "xlsx", "pptx", "image"})  # kinds that need the attachment bytes


class _Extracted(NamedTuple):
    """What one extractor found; ``extract_text`` adds the attachment's identity and hash."""

    status: AttachmentStatus
    text: str = ""
    page_count: int = 0
    needs_ocr_pages: tuple[int, ...] = ()
    reason: str = ""


def kind_for(filename: str, mime: str) -> AttachmentKind:
    """The attachment kind from the file suffix, falling back to the MIME type when the suffix says nothing."""
    kind = _SUFFIX_KINDS.get(os.path.splitext(filename)[1].lower())
    if kind is not None:
        return kind
    mime = mime.lower()
    if mime == "message/rfc822":
        return "eml"
    if mime == "application/vnd.ms-outlook":
        return "msg"
    if mime.startswith("image/"):
        return "image"
    if mime == "application/pdf":
        return "pdf"
    return "other"


def _extract_pdf(att: ParsedAttachment) -> _Extracted:
    pdf = pypdfium2.PdfDocument(att.data)
    try:
        page_count = len(pdf)
        texts: list[str] = []
        needs_ocr: list[int] = []
        for index in range(page_count):
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded().replace("\r\n", "\n")
                finally:
                    textpage.close()
            finally:
                page.close()
            if text.strip():
                texts.append(text)
            if len(text.strip()) < LOW_TEXT_CHARS_PER_PAGE and len(needs_ocr) < MAX_OCR_PAGES:
                needs_ocr.append(index)
    finally:
        pdf.close()
    return _Extracted("needs_ocr" if needs_ocr else "ok", "\n".join(texts), page_count, tuple(needs_ocr))


def _extract_docx(att: ParsedAttachment) -> _Extracted:
    doc = docx.Document(io.BytesIO(att.data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:  # a linked header repeats an earlier section's text
                lines.extend(p.text for p in part.paragraphs)
    return _Extracted("ok", "\n".join(line for line in lines if line.strip()))


def _xlsx_lines(wb: openpyxl.Workbook) -> tuple[list[str], bool]:
    """The sheet names and non-empty cell values as lines, and whether ``MAX_XLSX_CELLS`` cut them short."""
    lines: list[str] = []
    cells = 0
    for ws in wb.worksheets:
        lines.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            remaining = MAX_XLSX_CELLS - cells
            truncated = len(row) > remaining
            row = row[:remaining]
            cells += len(row)
            values = [str(v) for v in row if v is not None and str(v).strip()]
            if values:
                lines.append("\t".join(values))
            if truncated:
                return lines, True
    return lines, False


def _extract_xlsx(att: ParsedAttachment) -> _Extracted:
    wb = openpyxl.load_workbook(io.BytesIO(att.data), read_only=True, data_only=True)
    try:
        lines, capped = _xlsx_lines(wb)
    finally:
        wb.close()
    return _Extracted("partial" if capped else "ok", "\n".join(lines), reason="cell_cap" if capped else "")


def _pptx_shapes(shapes):
    """Every shape on a slide, looking inside groups."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _pptx_shapes(shape.shapes)
        else:
            yield shape


def _extract_pptx(att: ParsedAttachment) -> _Extracted:
    lines: list[str] = []
    for slide in pptx.Presentation(io.BytesIO(att.data)).slides:
        for shape in _pptx_shapes(slide.shapes):
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
            if shape.has_table:
                lines.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            lines.append(slide.notes_slide.notes_text_frame.text)
    return _Extracted("ok", "\n".join(line for line in lines if line.strip()))


def _extract_image(att: ParsedAttachment) -> _Extracted:
    return _Extracted("needs_ocr", page_count=1, needs_ocr_pages=(0,))


def _extract_message(att: ParsedAttachment) -> _Extracted:
    if att.nested is None:
        return _Extracted("unsupported", reason="no_nested_message")
    return _Extracted("ok", att.nested.subject + "\n" + att.nested.body_text)


_EXTRACTORS: dict[str, Callable[[ParsedAttachment], _Extracted]] = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
    "pptx": _extract_pptx,
    "image": _extract_image,
    "msg": _extract_message,
    "eml": _extract_message,
}


def _extract(att: ParsedAttachment) -> _Extracted:
    if att.kind == "image" and att.content_id is not None and att.size < INLINE_IMAGE_MIN_BYTES:
        return _Extracted("unsupported", reason="inline_image")
    extractor = _EXTRACTORS.get(att.kind)
    if extractor is None:
        return _Extracted("unsupported", reason="unsupported_type")
    if att.kind in _DATA_KINDS and att.data is None:
        return _Extracted("failed", reason="no_data")
    try:
        return extractor(att)
    except Exception as e:  # noqa: BLE001 - one damaged attachment must never fail its email
        if isinstance(e, pypdfium2.PdfiumError) and "password" in str(e).lower():
            return _Extracted("unsupported", reason="encrypted")
        return _Extracted("failed", reason=type(e).__name__)


def extract_text(att: ParsedAttachment) -> AttachmentText:
    """The searchable text of one attachment; never raises (see the module docstring)."""
    if att.size > MAX_ATTACHMENT_BYTES:  # not even hashed: the data stays unread
        found, sha256 = _Extracted("unsupported", reason="too_large"), None
    else:
        found = _extract(att)
        sha256 = hashlib.sha256(att.data).hexdigest() if isinstance(att.data, bytes) else None
    return AttachmentText(
        att.ordinal,
        att.filename,
        att.kind,
        att.size,
        found.status,
        found.text,
        found.page_count,
        list(found.needs_ocr_pages),
        sha256,
        found.reason,
    )


def render_pdf_page_png(data: bytes, page_index: int, dpi: int = 200) -> bytes:
    """One zero-based page of a PDF as PNG bytes, for the OCR step. Unlike ``extract_text`` this raises."""
    pdf = pypdfium2.PdfDocument(data)
    try:
        page = pdf[page_index]
        try:
            image = page.render(scale=dpi / 72).to_pil()
        finally:
            page.close()
    finally:
        pdf.close()
    buf = io.BytesIO()
    image.save(buf, "PNG")
    return buf.getvalue()
